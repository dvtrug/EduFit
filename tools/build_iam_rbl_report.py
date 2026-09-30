from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"E:\iykyk\S5\SWP391\Projects\EduFit\docs\reports\Bao_cao_RBL_Authentication_IAM_EduFit.docx")

NAVY = "17365D"
BLUE = "2F5597"
PALE = "EAF2F8"
GRAY = "F2F2F2"
BORDER = "D9D9D9"
GREEN = "E2F0D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), BORDER)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"), "true"); trPr.append(tblHeader)

def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True

def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table); set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = h; shade(c, NAVY); set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(font_size)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value); set_cell_margins(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2: shade(cells[i], "F7FAFC")
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(headers)-1) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: r.font.size = Pt(font_size)
        if str(row[-1]).upper() == "PASSED": shade(cells[-1], GREEN)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths): row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text); return p

def numbered(doc, text):
    p = doc.add_paragraph(style="List Number"); p.add_run(text); return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.12
for name, size in (("Title", 24), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    st = styles[name]; st.font.name = "Aptos Display"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    st.font.size = Pt(size); st.font.color.rgb = RGBColor(0,0,0); st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(5); st.paragraph_format.keep_with_next = True

# Cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(60)
r = p.add_run("TRƯỜNG ĐẠI HỌC  ______________________________")
r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0,0,0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MÔN HỌC SWP391  DỰ ÁN PHẦN MỀM"); r.bold = True; r.font.size = Pt(12)

p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(75)
p.add_run("BÁO CÁO KIỂM THỬ THEO YÊU CẦU")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RBL AUTHENTICATION TRONG MODULE IAM"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Dự án EduFit"); r.bold = True; r.font.size = Pt(14)

doc.add_paragraph()
meta = add_table(doc, ["Thông tin", "Nội dung"], [
    ("Nhóm thực hiện", "........................................................................"),
    ("Giảng viên", "........................................................................"),
    ("Phân hệ", "backend/modules/iam"),
    ("Phương pháp", "Requirements Based Logic và Test Driven Development"),
    ("Ngày xác minh", "30/09/2026"),
    ("Trạng thái", "32 test executions passed, 0 failures, 0 errors, 0 skipped"),
], widths=[4.2, 11.5], font_size=9.5)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(35)
p.add_run("Thành phố Hồ Chí Minh  2026").italic = True
doc.add_page_break()

# Executive summary and TOC
doc.add_heading("Tóm tắt báo cáo", level=1)
doc.add_paragraph(
    "Báo cáo trình bày business logic và bằng chứng kiểm thử tự động cho chức năng Authentication của module IAM trong dự án EduFit. "
    "Kết quả Maven Surefire được xác minh ngày 30/09/2026 cho thấy module IAM thực thi 32 lượt kiểm thử, toàn bộ đều thành công. "
    "Phạm vi bằng chứng hiện tại chủ yếu là unit test cho domain, application service và các thành phần bảo mật; báo cáo không xem đây là bằng chứng thay thế cho kiểm thử tích hợp PostgreSQL hoặc kiểm thử end to end."
)
add_table(doc, ["Chỉ số", "Kết quả"], [
    ("Test classes", "7"), ("Test executions", "32"), ("Passed", "32"),
    ("Failures", "0"), ("Errors", "0"), ("Skipped", "0"), ("Build", "SUCCESS")
], widths=[7.5, 7.5], font_size=9.5)

doc.add_heading("Mục lục", level=1)
for item in [
    "1  Mục tiêu và phạm vi", "2  Phương pháp và công nghệ kiểm thử",
    "3  Business rules và non functional requirements", "4  Business logic theo tính năng",
    "5  Ma trận truy vết RBL", "6  Kết quả thực thi", "7  Đánh giá và giới hạn", "8  Kết luận"
]:
    doc.add_paragraph(item)
doc.add_page_break()

doc.add_heading("1  Mục tiêu và phạm vi", level=1)
doc.add_heading("1.1  Mục tiêu", level=2)
doc.add_paragraph("Mục tiêu của báo cáo là mô tả quy tắc nghiệp vụ Authentication, chỉ ra phương thức hiện thực trong module IAM và cung cấp test case có thể chạy lại bằng JUnit 5. Mỗi kết luận PASSED phải truy vết được tới lớp kiểm thử và phương thức test tương ứng.")
doc.add_heading("1.2  Phạm vi đã kiểm thử", level=2)
for t in [
    "Thực thể Account, trạng thái, vai trò và lockout.", "PasswordPolicy và Argon2PasswordService.",
    "RegisterAccountService và AuthenticateAccountService.", "AccountToken và TokenCryptoService cho password reset.",
    "RequestPasswordResetService, ResetPasswordService và ChangePasswordService.",
    "IamFacade và AuthController đã tồn tại trong production code; tuy nhiên báo cáo này chỉ ghi nhận bằng chứng test hiện có."
]: bullet(doc,t)
doc.add_heading("1.3  Ngoài phạm vi bằng chứng", level=2)
doc.add_paragraph("Chưa có test tích hợp PostgreSQL/Testcontainers, test controller HTTP, test session/cookie hoặc test end to end trong tập 32 test hiện tại. Vì vậy các thuộc tính cơ sở dữ liệu và hành vi HTTP chưa được tuyên bố là đã được chứng minh bởi báo cáo này.")

doc.add_heading("2  Phương pháp và công nghệ kiểm thử", level=1)
add_table(doc, ["Hạng mục", "Áp dụng"], [
    ("Nền tảng", "Java 21 và Spring Boot 4.x"),
    ("Framework", "JUnit 5 Jupiter, Mockito, JUnit Assertions"),
    ("Loại test", "Unit tests cho domain và application services"),
    ("Thời gian", "Clock cố định hoặc Instant truyền vào, không dùng Thread.sleep"),
    ("Mật khẩu", "Argon2id với salt ngẫu nhiên"),
    ("Lệnh xác minh", "mvn -B -ntp -pl modules/iam -am test"),
], widths=[4.2, 11.5], font_size=9.2)

doc.add_heading("3  Business rules và non functional requirements", level=1)
rules = [
    ("BR-01", "Email", "Trim và lowercase trước khi tìm/lưu; từ chối email đã tồn tại."),
    ("BR-02", "Password policy", "Tối thiểu 8 ký tự, có chữ và số; chặn sử dụng lại mật khẩu."),
    ("BR-03", "Role", "Self-registration chỉ cho STUDENT, PARENT, TUTOR; cấm ADMIN."),
    ("BR-04", "Account status", "Đăng ký mới tạo account ACTIVE và có thể đăng nhập ngay."),
    ("BR-05", "One-time token", "Password-reset token có TTL 30 phút, dùng một lần, hết hạn tại đúng expiresAt."),
    ("BR-06", "Lockout", "Sai 5 lần liên tiếp khóa 15 phút; login thành công reset counter."),
    ("NFR-01", "Password storage", "Hash bằng Argon2id, salt ngẫu nhiên, không lưu plaintext."),
    ("NFR-03", "Token storage", "Token hash phải là SHA-256 hex 64 ký tự; không lưu raw token."),
    ("NFR-06", "Enumeration resistance", "Sai email và sai password trả cùng INVALID_CREDENTIALS; forgot password trả trung tính."),
    ("NFR-11", "Atomicity", "Reset password và consume token nằm trong transaction ở application service."),
]
add_table(doc, ["Mã", "Tên", "Business rule"], rules, widths=[2.0,4.0,10.0], font_size=8.8)

doc.add_heading("4  Business logic theo tính năng", level=1)
features = [
    ("Đăng ký", ["Chuẩn hóa email", "Kiểm tra trùng", "Kiểm tra role và password", "Băm Argon2id", "Tạo account ACTIVE và lưu"]),
    ("Đăng nhập", ["Chuẩn hóa email", "Dummy verify khi email không tồn tại", "Kiểm tra trạng thái và lockout", "Verify password", "Ghi failure hoặc success"]),
    ("Reset mật khẩu", ["Luôn phản hồi trung tính ở bước yêu cầu", "Sinh raw token và chỉ lưu SHA-256 hash", "Kiểm tra token và policy", "Đổi password hash", "Consume token trong transaction"]),
    ("Đổi mật khẩu", ["Verify mật khẩu hiện tại", "Chặn mật khẩu mới trùng cũ", "Áp dụng policy", "Băm và cập nhật password hash"]),
]
for name, steps in features:
    doc.add_heading(name, level=2)
    for s in steps: numbered(doc, s)

doc.add_heading("5  Ma trận truy vết RBL", level=1)
doc.add_paragraph("Bảng dưới đây liệt kê các test case được đặt mã rõ ràng trong mã nguồn. Một số phương thức parameterized tạo nhiều lượt thực thi, vì vậy số test executions của Surefire lớn hơn số dòng test case mô tả trong ma trận.")
matrix = [
 ("IAM-DOM-001","BR-04","AccountTest","shouldInitializeAccountInActiveStatus","PASSED"),
 ("IAM-DOM-002","BR-03","AccountTest","shouldPreventAdminRoleRegistration","PASSED"),
 ("IAM-DOM-005","BR-06","AccountTest","shouldIncrementFailedAttemptsAndLockAccountAtFifthFailure","PASSED"),
 ("IAM-DOM-006","BR-06","AccountTest","shouldDenyLoginWhenAccountIsCurrentlyLocked","PASSED"),
 ("IAM-DOM-007","BR-06","AccountTest","shouldAllowLoginAndResetCounterAfterLockExpires","PASSED"),
 ("IAM-DOM-008","BR-06","AccountTest","shouldResetFailedAttemptsOnSuccessfulLogin","PASSED"),
 ("IAM-DOM-009","BR-02","PasswordPolicyTest","shouldAcceptValidPassword","PASSED"),
 ("IAM-DOM-010","BR-02","PasswordPolicyTest","parameterized short or blank passwords","PASSED"),
 ("IAM-DOM-011","BR-02","PasswordPolicyTest","shouldRejectPasswordWithoutDigitsOrLetters","PASSED"),
 ("IAM-DOM-012","BR-05 NFR-03","AccountTokenTest","shouldCreateValidTokenWithCorrectExpiry","PASSED"),
 ("IAM-DOM-013","BR-05","AccountTokenTest","shouldConsumeTokenSuccessfullyOnce","PASSED"),
 ("IAM-DOM-014","BR-05","AccountTokenTest","shouldRejectExpiredOrAlreadyUsedToken","PASSED"),
 ("IAM-TOKEN-BOUNDARY-001","BR-05","AccountTokenTest","shouldConsiderTokenExpiredAtExactExpiryInstant","PASSED"),
 ("IAM-TOKEN-HASH-001","NFR-03","AccountTokenTest","shouldRejectTokenHashNotMeetingSha256HexRequirements","PASSED"),
 ("IAM-REG-001","BR-01 BR-03 BR-04","RegisterAccountServiceTest","shouldRegisterAccountSuccessfullyInActiveStatus","PASSED"),
 ("IAM-REG-003","BR-01","RegisterAccountServiceTest","shouldRejectRegistrationWhenEmailAlreadyExists","PASSED"),
 ("IAM-REG-004","BR-03","RegisterAccountServiceTest","shouldRejectAdminRegistrationAttempt","PASSED"),
 ("IAM-REG-005","BR-02","RegisterAccountServiceTest","shouldRejectRegistrationWithWeakPassword","PASSED"),
 ("IAM-AUTH-001","BR-06 NFR-01","AuthenticateAccountServiceTest","shouldAuthenticateSuccessfullyWhenCredentialsAreValid","PASSED"),
 ("IAM-AUTH-002","NFR-06","AuthenticateAccountServiceTest","shouldReturnSameErrorForWrongPasswordAndNonExistentEmail","PASSED"),
 ("IAM-AUTH-004","BR-06","AuthenticateAccountServiceTest","shouldLockAccountAfterFiveConsecutiveFailures","PASSED"),
 ("IAM-PWD-001","NFR-06","PasswordResetAndChangeTest","shouldReturnNeutralResultEvenWhenEmailDoesNotExist","PASSED"),
 ("IAM-PWD-002","BR-05 NFR-11","PasswordResetAndChangeTest","shouldResetPasswordSuccessfully","PASSED"),
 ("IAM-PWD-003","UC1.4","PasswordResetAndChangeTest","shouldRejectWhenCurrentPasswordIsIncorrect","PASSED"),
 ("IAM-PWD-004","BR-02","PasswordResetAndChangeTest","shouldRejectWhenNewPasswordMatchesCurrentPassword","PASSED"),
 ("IAM-SEC-001","NFR-01","Argon2PasswordServiceTest","shouldHashAndVerifyPasswordSuccessfully","PASSED"),
 ("IAM-SEC-002","NFR-01","Argon2PasswordServiceTest","shouldProduceDifferentHashesForSamePasswordDueToSalt","PASSED"),
 ("IAM-SEC-003","NFR-01","Argon2PasswordServiceTest","shouldRejectWrongPassword","PASSED"),
]
add_table(doc, ["Test case", "Rule", "Test class", "Test method", "Kết quả"], matrix, widths=[3.1,2.5,4.1,5.4,1.8], font_size=7.5)

doc.add_heading("6  Kết quả thực thi", level=1)
doc.add_paragraph("Lần chạy xác minh gần nhất được thực hiện tại 15:04:49 ngày 30/09/2026 theo múi giờ Asia Saigon.")
add_table(doc, ["Test class", "Tests", "Failures", "Errors", "Skipped"], [
 ("AccountTest",6,0,0,0),("AccountTokenTest",5,0,0,0),("Argon2PasswordServiceTest",3,0,0,0),
 ("AuthenticateAccountServiceTest",3,0,0,0),("PasswordPolicyTest",7,0,0,0),
 ("PasswordResetAndChangeTest",4,0,0,0),("RegisterAccountServiceTest",4,0,0,0),
 ("Tổng IAM",32,0,0,0)
], widths=[7.0,2.1,2.1,2.1,2.1], font_size=8.8)
doc.add_paragraph("Kết quả: BUILD SUCCESS. Thời gian module IAM: 6.427 giây. Tổng thời gian Maven reactor: 10.326 giây.")

doc.add_heading("7  Đánh giá và giới hạn", level=1)
doc.add_heading("7.1  Điểm đã chứng minh", level=2)
for t in ["Các rule BR-01 đến BR-06 đều có test tự động liên quan.", "Argon2id, salt ngẫu nhiên và xác minh mật khẩu được kiểm thử.", "Các trường hợp token hết hạn đúng deadline và hash sai định dạng được kiểm thử.", "Đăng nhập sai email và sai mật khẩu dùng cùng mã lỗi.", "Không có test thất bại, lỗi hoặc bị bỏ qua trong lần chạy xác minh."]: bullet(doc,t)
doc.add_heading("7.2  Giới hạn cần ghi nhận", level=2)
for t in ["Chưa có PostgreSQL/Testcontainers integration test trong tập test hiện tại.", "Annotation Transactional cho thấy ranh giới giao dịch nhưng unit test chưa chứng minh rollback thực tế của database.", "AuthController và session/cookie chưa có controller/integration test trong module IAM.", "Phạm vi Authentication không có bước xác minh email; tài khoản đăng ký hợp lệ được tạo trực tiếp ở trạng thái ACTIVE.", "Mockito phát cảnh báo dynamic agent trên JDK hiện tại; không làm test thất bại nhưng nên cấu hình agent trước khi nâng JDK."]: bullet(doc,t)

doc.add_heading("8  Kết luận", level=1)
doc.add_paragraph("Phần RBL Authentication hiện đã có bộ unit test chạy được và truy vết được tới các business rule chính. Lần xác minh ngày 30/09/2026 đạt 32 trên 32 lượt kiểm thử thành công. Báo cáo đủ dùng để trình bày phần business logic và phương thức test theo yêu cầu môn học, với điều kiện nhóm nêu rõ đây là bằng chứng unit test; kiểm thử tích hợp cơ sở dữ liệu và HTTP là hạng mục mở rộng tiếp theo.")

# Header/footer and page number field
for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = "EDUFIT  RBL AUTHENTICATION  IAM"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor(90,90,90)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Trang ")
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    run.font.size = Pt(8)

props = doc.core_properties
props.title = "Báo cáo RBL Authentication trong module IAM"
props.subject = "EduFit SWP391"
props.author = "Nhóm dự án EduFit"
props.keywords = "RBL, Authentication, IAM, JUnit 5, EduFit"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
